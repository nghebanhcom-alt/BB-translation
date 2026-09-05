from dataclasses import dataclass
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.worksheet.worksheet import Worksheet

# BR-GLOSS-05: column 1 = EN, column 2 = VI (or "(keep)"), column 3 = notes (optional).
_HEADER = ["EN", "VI", "Notes"]


class GlossaryExcelError(ValueError):
    """Raised when the uploaded Excel file is missing required columns."""


@dataclass
class GlossaryEntryData:
    term_en: str
    term_vi: str | None
    notes: str | None = None


def _select_data_sheet(workbook: Workbook) -> Worksheet:
    # workbook.active reflects whichever sheet was active when the file was
    # last saved in Excel/Google Sheets — a detail end users neither know nor
    # control, so a "Glossary" sheet placed after e.g. an instructions sheet
    # (as in data/glossary-starter.xlsx) would silently be skipped.
    for name in workbook.sheetnames:
        if name.strip().lower() == "glossary":
            return workbook[name]
    return workbook.active


def import_glossary_from_excel(file_path: Path) -> list[GlossaryEntryData]:
    workbook = load_workbook(file_path, read_only=True, data_only=True)
    sheet = _select_data_sheet(workbook)

    rows = sheet.iter_rows(values_only=True)
    try:
        header = next(rows)
    except StopIteration as exc:
        raise GlossaryExcelError("File Excel rong, khong co header row") from exc

    if header is None or len(header) < 2 or not header[0] or not header[1]:
        raise GlossaryExcelError(
            "Thieu cot bat buoc: can cot A (EN) va cot B (VI). "
            f"Header nhan duoc: {header}"
        )

    entries: list[GlossaryEntryData] = []
    for row in rows:
        if row is None or all(cell is None for cell in row):
            continue

        term_en = str(row[0]).strip() if len(row) > 0 and row[0] is not None else ""
        if not term_en:
            continue

        term_vi_raw = row[1] if len(row) > 1 and row[1] is not None else None
        term_vi = str(term_vi_raw).strip() if term_vi_raw is not None else None
        term_vi = term_vi if term_vi else None

        notes_raw = row[2] if len(row) > 2 and row[2] is not None else None
        notes = str(notes_raw).strip() if notes_raw is not None else None
        notes = notes if notes else None

        entries.append(GlossaryEntryData(term_en=term_en, term_vi=term_vi, notes=notes))

    workbook.close()
    return entries


def export_glossary_to_excel(entries: list[GlossaryEntryData], output_path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Glossary"

    sheet.append(_HEADER)
    for cell in sheet[1]:
        cell.font = Font(bold=True)

    for entry in entries:
        sheet.append([entry.term_en, entry.term_vi or "", entry.notes or ""])

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
