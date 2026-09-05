from pathlib import Path

import pytest
from openpyxl import Workbook

from src.utils.excel_utils import (
    GlossaryEntryData,
    GlossaryExcelError,
    export_glossary_to_excel,
    import_glossary_from_excel,
)


def test_export_then_import_roundtrip(tmp_path: Path) -> None:
    entries = [
        GlossaryEntryData(term_en="ganache", term_vi="(keep)", notes="chocolate ganache"),
        GlossaryEntryData(term_en="buttercream", term_vi="kem phu bo", notes=None),
        GlossaryEntryData(term_en="proof", term_vi="u bot", notes="fermentation step"),
    ]
    output_path = tmp_path / "glossary.xlsx"

    export_glossary_to_excel(entries, output_path)
    imported = import_glossary_from_excel(output_path)

    assert len(imported) == len(entries)
    for original, roundtripped in zip(entries, imported, strict=True):
        assert roundtripped.term_en == original.term_en
        assert roundtripped.term_vi == original.term_vi
        assert roundtripped.notes == original.notes


def test_import_skips_empty_rows_and_trims_whitespace(tmp_path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["EN", "VI", "Notes"])
    sheet.append(["  fondant  ", "  (keep)  ", None])
    sheet.append([None, None, None])
    sheet.append(["crumb coat", "lop kem lot", "thin base layer"])
    path = tmp_path / "with_gaps.xlsx"
    workbook.save(path)

    entries = import_glossary_from_excel(path)

    assert len(entries) == 2
    assert entries[0].term_en == "fondant"
    assert entries[0].term_vi == "(keep)"
    assert entries[1].term_en == "crumb coat"
    assert entries[1].notes == "thin base layer"


def test_import_missing_required_column_raises(tmp_path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Term"])
    sheet.append(["fondant"])
    path = tmp_path / "bad_header.xlsx"
    workbook.save(path)

    with pytest.raises(GlossaryExcelError):
        import_glossary_from_excel(path)


def test_import_empty_workbook_raises(tmp_path: Path) -> None:
    workbook = Workbook()
    path = tmp_path / "empty.xlsx"
    workbook.save(path)

    with pytest.raises(GlossaryExcelError):
        import_glossary_from_excel(path)


def test_import_finds_glossary_sheet_when_not_active(tmp_path: Path) -> None:
    """Regression: data/glossary-starter.xlsx has an instructions sheet saved
    as active, in front of the "Glossary" data sheet. workbook.active alone
    reads the wrong sheet and fails with a misleading "missing column" error
    even though the file is valid.
    """
    workbook = Workbook()
    instructions = workbook.active
    instructions.title = "Huong dan"
    instructions.append(["Day la file huong dan, khong phai du lieu"])

    glossary = workbook.create_sheet("Glossary")
    glossary.append(["EN", "VI", "Notes"])
    glossary.append(["fondant", "(keep)", None])
    glossary.append(["ganache", "kem ganache", "chocolate ganache"])

    workbook.active = 0  # "Huong dan" saved as the active sheet, as Excel would
    path = tmp_path / "multi_sheet.xlsx"
    workbook.save(path)

    entries = import_glossary_from_excel(path)

    assert len(entries) == 2
    assert entries[0].term_en == "fondant"
    assert entries[1].term_en == "ganache"
    assert entries[1].notes == "chocolate ganache"
